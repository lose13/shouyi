from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2] / '守忆'
SOURCE = ROOT / 'app' / '策略说明.md'
TARGET = ROOT / '策略说明报告.docx'

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2.3), Cm(2.3)
section.left_margin, section.right_margin = Cm(2.5), Cm(2.5)

for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attr), 'Microsoft YaHei')
    style.paragraph_format.widow_control = True

normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.45
normal.paragraph_format.space_after = Pt(8)
title = doc.styles['Title']
title.font.size, title.font.bold = Pt(23), True
title.paragraph_format.space_after = Pt(12)
heading = doc.styles['Heading 1']
heading.font.size, heading.font.bold = Pt(14), True
heading.paragraph_format.space_before = Pt(16)
heading.paragraph_format.space_after = Pt(8)
heading.paragraph_format.keep_with_next = True
subtitle = doc.styles['Subtitle']
subtitle.font.size = Pt(11)
subtitle.font.italic = False
subtitle.paragraph_format.space_after = Pt(18)

source = SOURCE.read_text(encoding='utf-8')
for line in source.splitlines():
    line = line.strip()
    if not line:
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], style='Title')
        doc.add_paragraph('赛题三  基于角色隔离的记忆管理', style='Subtitle')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], style='Heading 1')
    else:
        doc.add_paragraph(line)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)
doc.core_properties.title = '守忆角色记忆管理策略说明报告'
doc.core_properties.subject = '赛题三 基于角色隔离的记忆管理'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.save(TARGET)
print('DOCX report created.')
